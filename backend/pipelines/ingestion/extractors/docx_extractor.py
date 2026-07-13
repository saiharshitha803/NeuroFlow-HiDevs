from docx import Document

from .base import ExtractedPage


class DOCXExtractor:
    """
    Extract text, headings, headers and tables from DOCX files.
    """

    def extract(
        self,
        file_path: str,
    ) -> list[ExtractedPage]:

        document = Document(file_path)

        pages = []

        page_number = 1

        # Paragraphs
        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if not text:
                continue

            level = "paragraph"

            if paragraph.style.name.startswith("Heading"):

                level = paragraph.style.name.lower()

            pages.append(
                ExtractedPage(
                    page_number=page_number,
                    content=text,
                    content_type="text",
                    metadata={
                        "level": level,
                    },
                )
            )

        # Tables
        for table in document.tables:

            rows = []

            for row in table.rows:
                rows.append(
                    " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                    )
                )

            pages.append(
                ExtractedPage(
                    page_number=page_number,
                    content="\n".join(rows),
                    content_type="table",
                    metadata={},
                )
            )

        # Headers
        for section in document.sections:

            header = section.header

            for paragraph in header.paragraphs:

                if paragraph.text.strip():

                    pages.append(
                        ExtractedPage(
                            page_number=page_number,
                            content=paragraph.text.strip(),
                            content_type="text",
                            metadata={
                                "section": "header",
                            },
                        )
                    )

        return pages