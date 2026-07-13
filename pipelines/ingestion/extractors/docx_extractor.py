from pathlib import Path

from docx import Document

from pipelines.ingestion.models import ExtractedPage


class DOCXExtractor:
    """
    Extract text, headings, tables, and headers
    from Microsoft Word (.docx) documents.
    """

    def extract(
        self,
        file_path: str | Path,
    ) -> list[ExtractedPage]:

        document = Document(str(file_path))

        extracted_pages: list[ExtractedPage] = []

        page_number = 1

        # ----------------------------
        # Extract paragraphs & headings
        # ----------------------------

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if not text:
                continue

            style_name = paragraph.style.name

            metadata = {}

            if style_name.startswith("Heading"):

                level = style_name.replace(
                    "Heading",
                    ""
                ).strip()

                metadata = {
                    "level": f"h{level}",
                    "section": text,
                }

            extracted_pages.append(

                ExtractedPage(

                    page_number=page_number,

                    content=text,

                    content_type="text",

                    metadata=metadata,

                )

            )

        # ----------------------------
        # Extract tables
        # ----------------------------

        for table in document.tables:

            rows = []

            for row in table.rows:

                cells = [

                    cell.text.strip()

                    for cell in row.cells

                ]

                rows.append(" | ".join(cells))

            extracted_pages.append(

                ExtractedPage(

                    page_number=page_number,

                    content="\n".join(rows),

                    content_type="table",

                    metadata={
                        "source": str(file_path),
                    },

                )

            )

        # ----------------------------
        # Extract headers
        # ----------------------------

        for section in document.sections:

            header = section.header

            for paragraph in header.paragraphs:

                text = paragraph.text.strip()

                if not text:
                    continue

                extracted_pages.append(

                    ExtractedPage(

                        page_number=page_number,

                        content=text,

                        content_type="text",

                        metadata={
                            "header": True,
                        },

                    )

                )

        return extracted_pages